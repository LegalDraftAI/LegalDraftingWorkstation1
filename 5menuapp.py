import streamlit as st
import os
import requests
import pandas as pd
from datetime import datetime
from google import genai
from dotenv import load_dotenv
from docx import Document
from fpdf import FPDF
import io

# 1. SETUP & PERSISTENCE
load_dotenv()
INITIAL_G_KEY = os.getenv("GOOGLE_API_KEY", "").strip()
INITIAL_IK_TOKEN = os.getenv("INDIAN_KANOON_TOKEN", "").strip()
HISTORY_FILE = "kerala_legal_history.csv"

st.set_page_config(page_title="Kerala Senior Advocate Workstation", layout="wide")

# Initialize Session State
if "google_key" not in st.session_state: st.session_state.google_key = INITIAL_G_KEY
if "ik_token" not in st.session_state: st.session_state.ik_token = INITIAL_IK_TOKEN
if "final_master" not in st.session_state: st.session_state.final_master = ""
if "search_results" not in st.session_state: st.session_state.search_results = []
if "draft_done" not in st.session_state: st.session_state.draft_done = False

def reset_workstation():
    st.session_state.final_master = ""
    st.session_state.search_results = []
    st.session_state.draft_done = False

KERALA_DISTRICTS = ["Alappuzha", "Ernakulam", "Idukki", "Kannur", "Kasaragod", "Kollam", "Kottayam", "Kozhikode", "Malappuram", "Palakkad", "Pathanamthitta", "Thiruvananthapuram", "Thrissur", "Wayanad"]

# 2. EXPORT ENGINES
def generate_docx(text):
    doc = Document()
    doc.add_paragraph(text)
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()

def generate_pdf(text):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=11)
    safe_text = text.encode('latin-1', 'replace').decode('latin-1')
    pdf.multi_cell(0, 8, txt=safe_text)
    return pdf.output(dest='S').encode('latin-1')

# 3. INTERFACE
tab_work, tab_settings = st.tabs(["⚖️ Workstation", "⚙️ Settings"])

with tab_settings:
    st.header("🔑 API Configuration")
    st.session_state.google_key = st.text_input("Gemini API Key", value=st.session_state.google_key, type="password")
    st.session_state.ik_token = st.text_input("Indian Kanoon Token", value=st.session_state.ik_token, type="password")
    if st.button("💾 Save Settings"): st.rerun()

with tab_work:
    with st.sidebar:
        st.title("👨‍⚖️ Kerala Admin")
        doc_type = st.selectbox("Petition Type", ["Bail Application", "Cheque Bounce (NI Act)", "Legal Notice", "Maintenance Petition", "Divorce Petition", "Writ Petition"])
        jurisdiction = st.selectbox("Court Level", ["High Court", "District Court", "Family Court"])
        
        # Kerala Logic: Lock Ernakulam for High Court
        is_hc = (jurisdiction == "High Court")
        selected_district = st.selectbox(
            "Select District", 
            KERALA_DISTRICTS, 
            index=KERALA_DISTRICTS.index("Ernakulam") if is_hc else 0,
            disabled=is_hc
        )
        
        st.divider()
        if st.button("🧹 Clear Workstation", use_container_width=True):
            reset_workstation()
            st.rerun()

    # Success Celebration (Triggers after rerun)
    if st.session_state.draft_done:
        st.balloons()
        st.success("✅ Drafting Complete! The Senior Advocate has prepared the template below.", icon="⚖️")
        st.session_state.draft_done = False

    st.title(f"Drafting: {doc_type}")
    facts = st.text_area("Case Facts:", placeholder="Briefly enter the details (e.g., date of cheque, notice details)...", height=150)

    # ACTION BUTTONS
    c1, c2, _ = st.columns([1, 1, 3])
    
    with c1:
        if st.button("🔍 Search Precedents", use_container_width=True):
            if len(st.session_state.ik_token) < 5:
                st.session_state.search_results = [{"title": f"Kerala Ruling on {doc_type}", "headline": "A summary of relevant Kerala precedents for this matter."}]
            else:
                try:
                    headers = {'Authorization': f'Token {st.session_state.ik_token}'}
                    r = requests.post(f"https://api.indiankanoon.org/search/?formInput={facts}", headers=headers)
                    if r.status_code == 200: st.session_state.search_results = r.json().get('docs', [])
                except: st.error("Search Error.")
            st.rerun()

    with c2:
        if st.button("🚀 Draft Petition", type="primary", use_container_width=True):
            if not st.session_state.google_key:
                st.error("Missing Gemini Key!")
            else:
                # 4. TIMER & GENERATION
                with st.spinner("Senior Advocate is drafting... please wait.", show_time=True):
                    client = genai.Client(api_key=st.session_state.google_key)
                    court_label = "THE HIGH COURT OF KERALA AT ERNAKULAM" if is_hc else f"THE {jurisdiction.upper()} AT {selected_district.upper()}, KERALA"
                    
                    full_prompt = f"""
                    Act as a Senior Advocate of the Kerala High Court. 
                    Draft a formal {doc_type} for {court_label}.
                    RULES: Refer to Petitioner as "PARTY A" and Respondent as "PARTY B".
                    DO NOT invent specific names or addresses.
                    FACTS: {facts}
                    """
                    res = client.models.generate_content(model='gemini-2.5-flash', contents=full_prompt)
                    st.session_state.final_master = res.text
                    st.session_state.draft_done = True
                    st.rerun()

    # PRECEDENTS DISPLAY
    if st.session_state.search_results:
        with st.expander("📚 Relevant Kerala Cases", expanded=True):
            for doc in st.session_state.search_results[:3]:
                st.markdown(f"**{doc.get('title')}**")
                st.caption(doc.get('headline'))
                st.divider()

    # FINAL EDITOR
    if st.session_state.final_master:
        st.divider()
        st.subheader("📜 Live Template Editor")
        st.info("💡 Edit the placeholders (PARTY A, PARTY B) with client details below.")
        edited_text = st.text_area("Final Document:", value=st.session_state.final_master, height=500)
        
        ec1, ec2, ec3 = st.columns(3)
        with ec1:
            if st.button("💾 Save to Local History"):
                pd.DataFrame([{"Date": datetime.now(), "Type": doc_type, "Draft": edited_text}]).to_csv(HISTORY_FILE, mode='a', header=not os.path.exists(HISTORY_FILE), index=False)
                st.toast("Saved to History!")
        with ec2:
            st.download_button("📥 Word (.docx)", data=generate_docx(edited_text), file_name=f"Kerala_{doc_type}.docx")
        with ec3:
            st.download_button("📥 PDF (.pdf)", data=generate_pdf(edited_text), file_name=f"Kerala_{doc_type}.pdf")